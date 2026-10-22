"""Integration tests: ingestion pipeline against real PostgreSQL."""

from __future__ import annotations

import pytest

from tests.conftest import requires_db

pytestmark = [pytest.mark.integration, requires_db]


@requires_db
class TestIngestionPerFormat:
    def test_markdown_ingestion_preserves_sections(self, db_session, acme_user):
        from app.ingestion.pipeline import run_ingestion

        payload = (
            "# Leave Policy\n\nEmployees receive 22 days of paid annual leave.\n\n"
            "## Sick Leave\n\nA medical certificate is required after three days.\n"
        ).encode()
        result = run_ingestion(db_session, filename="leave_policy.md", payload=payload,
                               tenant_id=acme_user.tenant_id, user_id=acme_user.id)
        db_session.commit()
        assert result.status == "completed"
        assert result.chunk_count >= 2

        from sqlalchemy import select

        from app.models import Document, DocumentChunk

        document = db_session.get(Document, result.document_id)
        assert document.status == "completed"
        chunks = list(db_session.scalars(
            select(DocumentChunk).where(DocumentChunk.document_id == document.id)
        ))
        sections = {c.section for c in chunks}
        assert "Leave Policy" in sections
        assert "Sick Leave" in sections
        assert all(c.embedding is not None for c in chunks)  # embeddings stored

    def test_txt_ingestion(self, db_session, acme_user):
        from app.ingestion.pipeline import run_ingestion

        payload = b"Benefits summary.\n\nThe wellness benefit reimburses 300 USD per year."
        result = run_ingestion(db_session, filename="benefits.txt", payload=payload,
                               tenant_id=acme_user.tenant_id)
        db_session.commit()
        assert result.status == "completed"
        assert result.chunk_count >= 1

    def test_html_ingestion_preserves_headings(self, db_session, acme_user):
        from app.ingestion.pipeline import run_ingestion

        payload = (
            "<html><body><h1>Travel Policy</h1><p>Book via the portal.</p>"
            "<h2>Ground Transport</h2><p>Taxis are reimbursable.</p></body></html>"
        ).encode()
        result = run_ingestion(db_session, filename="travel.html", payload=payload,
                               tenant_id=acme_user.tenant_id)
        db_session.commit()
        assert result.status == "completed"

        from sqlalchemy import select

        from app.models import DocumentChunk

        chunks = list(db_session.scalars(select(DocumentChunk).where(
            DocumentChunk.document_id == result.document_id)))
        sections = {c.section for c in chunks}
        assert "Travel Policy" in sections and "Ground Transport" in sections

    def test_docx_ingestion(self, db_session, acme_user, tmp_path):
        import docx as docx_lib

        from app.ingestion.pipeline import run_ingestion

        document = docx_lib.Document()
        document.add_heading("Onboarding Guide", 0)
        document.add_heading("Day One", 1)
        document.add_paragraph("Receive laptop, badge, and hardware security key from IT.")
        path = tmp_path / "onboarding.docx"
        document.save(str(path))

        result = run_ingestion(db_session, filename="onboarding.docx", payload=path.read_bytes(),
                               tenant_id=acme_user.tenant_id)
        db_session.commit()
        assert result.status == "completed"
        assert result.chunk_count >= 1

    def test_pdf_ingestion_preserves_pages(self, db_session, acme_user, tmp_path):
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate

        from app.ingestion.pipeline import run_ingestion

        path = tmp_path / "vendor.pdf"
        doc = SimpleDocTemplate(str(path), pagesize=letter)
        styles = getSampleStyleSheet()
        doc.build([
            Paragraph("Vendor Guidelines", styles["Title"]),
            Paragraph("Payment terms are net 45 days from invoice receipt.", styles["BodyText"]),
            PageBreak(),
            Paragraph("Annual Review", styles["Heading1"]),
            Paragraph("Active vendors are reviewed annually for performance.", styles["BodyText"]),
        ])
        result = run_ingestion(db_session, filename="vendor.pdf", payload=path.read_bytes(),
                               tenant_id=acme_user.tenant_id)
        db_session.commit()
        assert result.status == "completed"

        from sqlalchemy import select

        from app.models import DocumentChunk

        chunks = list(db_session.scalars(select(DocumentChunk).where(
            DocumentChunk.document_id == result.document_id)))
        pages = {c.page_number for c in chunks}
        assert 1 in pages and 2 in pages          # page numbers preserved
        assert all(p is not None for p in pages)  # every chunk carries a page

    def test_full_sample_corpus(self, ingested_corpus):
        assert len(ingested_corpus) == 10
        assert all(r.status == "completed" for r in ingested_corpus)
        assert sum(r.chunk_count for r in ingested_corpus) > 30


@requires_db
class TestIngestionFailures:
    def test_empty_document_records_failed_job(self, db_session, acme_user):
        from app.ingestion.pipeline import run_ingestion
        from app.models import Document

        result = run_ingestion(db_session, filename="empty.md", payload=b"# only heading\n\n",
                               tenant_id=acme_user.tenant_id)
        db_session.commit()
        # heading-only markdown still yields one chunk; use truly empty txt
        result = run_ingestion(db_session, filename="empty.txt", payload=b"   ",
                               tenant_id=acme_user.tenant_id)
        db_session.commit()
        assert result.status == "failed"
        assert result.error
        document = db_session.get(Document, result.document_id)
        assert document.status == "failed"
        assert document.error_message

    def test_malformed_pdf_records_failure(self, db_session, acme_user):
        from app.ingestion.pipeline import run_ingestion

        result = run_ingestion(db_session, filename="broken.pdf",
                               payload=b"%PDF-1.4 totally not a real pdf body",
                               tenant_id=acme_user.tenant_id)
        db_session.commit()
        assert result.status == "failed"
        assert result.error

    def test_unsupported_type_rejected_before_db_write(self, db_session, acme_user):
        import uuid

        from app.core.errors import UnsupportedFileType
        from app.ingestion.pipeline import run_ingestion
        from app.models import Document

        before = len(list(db_session.query(Document)))
        with pytest.raises(UnsupportedFileType):
            run_ingestion(db_session, filename="virus.exe", payload=b"MZ...",
                          tenant_id=acme_user.tenant_id)
        db_session.rollback()
        after = len(list(db_session.query(Document)))
        assert after == before  # no partial rows

    def test_scanned_image_pdf_reports_no_text(self, db_session, acme_user):
        from app.ingestion.pipeline import run_ingestion

        # a real but empty-content PDF page (pypdf extracts no text)
        from reportlab.lib.pagesizes import letter
        from reportlab.pdfgen import canvas as pdf_canvas
        import io

        buffer = io.BytesIO()
        canvas = pdf_canvas.Canvas(buffer, pagesize=letter)
        canvas.showPage()
        canvas.save()
        result = run_ingestion(db_session, filename="blank.pdf", payload=buffer.getvalue(),
                               tenant_id=acme_user.tenant_id)
        db_session.commit()
        assert result.status == "failed"
        assert "no extractable text" in (result.error or "")

    def test_ingestion_timing_recorded(self, db_session, acme_user):
        from app.ingestion.pipeline import run_ingestion

        payload = b"# Policy\n\n" + b"Some substantive policy text about remote work. " * 20
        result = run_ingestion(db_session, filename="timing.md", payload=payload,
                               tenant_id=acme_user.tenant_id)
        db_session.commit()
        assert result.status == "completed"
        assert "parse_ms" in result.timing_ms
        assert "embedding_ms" in result.timing_ms
        assert result.duration_ms > 0
