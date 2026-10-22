"""Unit tests: parsers, doc-type detection, filename safety."""

from __future__ import annotations

import pytest

from app.core.errors import UnsupportedFileType
from app.ingestion.parsers import get_parser
from app.ingestion.parsers.base import detect_doc_type, safe_filename


class TestDetectDocType:
    def test_extensions(self):
        assert detect_doc_type("handbook.pdf", b"%PDF-1.7 ...") == "pdf"
        assert detect_doc_type("guide.DOCX", b"PK\x03\x04 rest") == "docx"
        assert detect_doc_type("notes.txt", b"hello") == "txt"
        assert detect_doc_type("policy.md", b"# hi") == "md"
        assert detect_doc_type("policy.markdown", b"# hi") == "md"
        assert detect_doc_type("page.html", b"<html></html>") == "html"
        assert detect_doc_type("page.htm", b"<html></html>") == "html"

    def test_unsupported_extension_rejected(self):
        with pytest.raises(UnsupportedFileType):
            detect_doc_type("malware.exe", b"MZ...")
        with pytest.raises(UnsupportedFileType):
            detect_doc_type("archive.zip", b"PK\x03\x04")

    def test_pdf_magic_bytes_enforced(self):
        with pytest.raises(UnsupportedFileType):
            detect_doc_type("fake.pdf", b"just text, not a pdf")

    def test_docx_magic_bytes_enforced(self):
        with pytest.raises(UnsupportedFileType):
            detect_doc_type("fake.docx", b"plain text")


class TestSafeFilename:
    def test_strips_paths(self):
        assert safe_filename("../../etc/passwd") == "passwd"
        assert safe_filename("C:\\Users\\x\\secret.pdf") == "secret.pdf"

    def test_strips_special_chars(self):
        assert safe_filename("my<>|:file*.pdf") == "myfile.pdf"

    def test_limits_length(self):
        assert len(safe_filename("a" * 5000 + ".pdf")) <= 200

    def test_empty_becomes_placeholder(self):
        assert safe_filename("///") == "uploaded_file"


class TestMarkdownParser:
    def test_sections_and_lists(self):
        payload = (
            "# Security Policy\n\n"
            "Passwords must rotate every 90 days.\n\n"
            "## Rules\n\n"
            "- Use a password manager\n"
            "- Never share credentials\n"
        ).encode()
        parsed = get_parser("md").parse("security_policy.md", payload)
        sections = [b.section for b in parsed.blocks if b.kind == "paragraph" or b.kind == "list_item"]
        assert "Security Policy" in sections
        assert "Rules" in sections
        list_items = [b for b in parsed.blocks if b.kind == "list_item"]
        assert any("password manager" in b.text for b in list_items)
        assert parsed.title == "Security Policy"

    def test_empty_raises(self):
        from app.ingestion.parsers.base import ParseError

        with pytest.raises(ParseError):
            get_parser("md").parse("empty.md", b"   \n\n  ")


class TestHtmlParser:
    def test_headings_and_scripts_excluded(self):
        payload = (
            "<html><head><title>Travel Policy</title>"
            "<script>alert('evil')</script></head>"
            "<body><h1>Travel Policy</h1><p>Book through the portal.</p>"
            "<h2>Ground Transport</h2><p>Taxis are reimbursable.</p>"
            "</body></html>"
        ).encode()
        parsed = get_parser("html").parse("travel_policy.html", payload)
        full = " ".join(b.text for b in parsed.blocks)
        assert "evil" not in full  # script content must never leak
        assert parsed.title == "Travel Policy"
        sections = {b.section for b in parsed.blocks}
        assert "Travel Policy" in sections and "Ground Transport" in sections


class TestTxtParser:
    def test_paragraphs(self):
        payload = b"First paragraph.\n\nSecond paragraph about expenses."
        parsed = get_parser("txt").parse("notes.txt", payload)
        assert len(parsed.blocks) == 2
        assert parsed.blocks[1].text.startswith("Second paragraph")

    def test_empty_raises(self):
        from app.ingestion.parsers.base import ParseError

        with pytest.raises(ParseError):
            get_parser("txt").parse("empty.txt", b"")


class TestPdfParser:
    def test_malformed_pdf_raises_parse_error(self):
        from app.ingestion.parsers.base import ParseError

        # valid magic, garbage body -> ParseError (recorded as failed ingestion)
        with pytest.raises(ParseError):
            get_parser("pdf").parse("broken.pdf", b"%PDF-1.4 this is not really a pdf")


class TestDocxParser:
    def test_generated_docx_roundtrip(self, tmp_path):
        import docx as docx_lib

        document = docx_lib.Document()
        document.add_heading("Onboarding", 0)
        document.add_heading("Day One", level=1)
        document.add_paragraph("Receive laptop and badge from IT.")
        path = tmp_path / "onboarding.docx"
        document.save(str(path))

        parsed = get_parser("docx").parse("onboarding.docx", path.read_bytes())
        assert parsed.title == "Onboarding"
        day_one = [b for b in parsed.blocks if b.section == "Day One"]
        assert any("laptop" in b.text for b in day_one)
