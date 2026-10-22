"""Generate the binary sample documents (DOCX + PDF) for the corpus.

Run:  python scripts/generate_sample_docs.py
Writes: data/sample_documents/onboarding_guide.docx, vendor_guidelines.pdf
"""

from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "data" / "sample_documents"


def build_docx() -> None:
    import docx

    doc = docx.Document()
    doc.add_heading("Onboarding Guide", level=0)
    doc.add_paragraph(
        "This guide walks new Acme Corp employees through their first 90 days."
    )

    doc.add_heading("Day One", level=1)
    doc.add_paragraph(
        "New employees receive their laptop, badge, and hardware security key "
        "from IT. Accounts for email, the HR portal, and the corporate drive "
        "are created before arrival. The IT team walks every new joiner "
        "through multi-factor authentication enrolment on day one."
    )

    doc.add_heading("First Week", level=1)
    doc.add_paragraph(
        "During the first week, new employees complete mandatory training: "
        "Security Awareness (2 hours), Code of Conduct (1 hour), and Data "
        "Protection Basics (1 hour). Manager one-to-ones happen twice in the "
        "first week, and a buddy from the team is assigned for the first "
        "month."
    )

    doc.add_heading("Probation Period", level=1)
    doc.add_paragraph(
        "The probation period lasts 90 days. A structured review happens at "
        "day 30, day 60, and day 90 with written feedback. Remote work "
        "eligibility begins after successful completion of probation, as "
        "defined in the Remote Work Policy."
    )

    doc.add_heading("First 90 Days Checklist", level=1)
    doc.add_paragraph("Complete security and conduct training", style="List Bullet")
    doc.add_paragraph("Set up the password manager and enrol MFA devices", style="List Bullet")
    doc.add_paragraph("Meet all direct stakeholders introduced by your manager", style="List Bullet")
    doc.add_paragraph("Agree on goals for the first quarter", style="List Bullet")

    doc.add_heading("Escalation and Help", level=1)
    doc.add_paragraph(
        "For HR questions contact People Operations through the HR portal. "
        "For equipment or access problems use the IT helpdesk. Security "
        "incidents follow the reporting duties in the Security Policy and "
        "must be reported within 24 hours."
    )

    OUT.mkdir(parents=True, exist_ok=True)
    doc.save(OUT / "onboarding_guide.docx")


def build_pdf() -> None:
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    doc = SimpleDocTemplate(
        str(OUT / "vendor_guidelines.pdf"),
        pagesize=letter,
        title="Vendor Guidelines",
        author="Acme Corp Procurement",
    )
    styles = getSampleStyleSheet()
    story = [
        Paragraph("Vendor Guidelines", styles["Title"]),
        Spacer(1, 0.2 * inch),
        Paragraph(
            "These guidelines define how Acme Corp selects, onboards, and "
            "manages vendors. They apply to all contracts above 5,000 USD "
            "in annual value.", styles["BodyText"],
        ),
        Paragraph("Vendor Registration", styles["Heading1"]),
        Paragraph(
            "Prospective vendors must register in the procurement portal, "
            "provide a W-9 or equivalent tax form, certificates of insurance, "
            "and two customer references. Registration review takes up to 15 "
            "working days.", styles["BodyText"],
        ),
        Paragraph("Security Assessment", styles["Heading1"]),
        Paragraph(
            "All vendors handling Confidential or Restricted data must "
            "complete a security questionnaire and pass a third-party risk "
            "assessment before contract signature. Vendors with system access "
            "must sign the Data Processing Addendum and enforce multi-factor "
            "authentication for all named accounts.", styles["BodyText"],
        ),
        Paragraph("Payment Terms", styles["Heading1"]),
        Paragraph(
            "Standard payment terms are net 45 days from invoice receipt. "
            "Invoices must reference a purchase order number; invoices "
            "without one are returned. Early-payment discounts of 2 percent "
            "or more are automatically accepted by Accounts Payable.",
            styles["BodyText"],
        ),
        Paragraph("Gifts and Hospitality", styles["Heading1"]),
        Paragraph(
            "Vendors may not offer gifts above 50 USD to any employee. "
            "Invitations to vendor-hosted events must be approved by "
            "Compliance in advance. This mirrors the employee obligations in "
            "the Employee Handbook.", styles["BodyText"],
        ),
        Paragraph("Annual Review", styles["Heading1"]),
        Paragraph(
            "Active vendors are reviewed annually for performance, security "
            "posture, and contract compliance. Vendors failing two "
            "consecutive reviews enter a remediation plan or are retired.",
            styles["BodyText"],
        ),
    ]
    doc.build(story)


if __name__ == "__main__":
    try:
        build_docx()
        print("wrote onboarding_guide.docx")
    except Exception as exc:
        print(f"docx generation failed: {exc}", file=sys.stderr)
        raise
    build_pdf()
    print("wrote vendor_guidelines.pdf")
